import json
import math
from pathlib import Path
import shutil
import struct
import subprocess
import sys
import unittest
from unittest.mock import patch
import uuid
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from m31tool import formats


def settings(**changes):
    result = {
        "transport": "TCP", "host": "", "tcp_port": 502,
        "serial_port": "", "baudrate": 9600, "parity": "N",
        "stopbits": 1, "unit_id": 1, "timeout": 2.0,
        "object_type": "Holding registers", "address": 0, "count": 1,
    }
    result.update(changes)
    return result


class LocalFilesTestCase(unittest.TestCase):
    def setUp(self):
        self.directory = Path(__file__).parent / (".formats-test-" + uuid.uuid4().hex)
        self.directory.mkdir()
        self.addCleanup(shutil.rmtree, self.directory)
        self.xlsx = self.directory / "values.xlsx"
        self.profile = self.directory / "profile.json"

    def workbook(self, rows, extra_sheet=False):
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        if extra_sheet:
            workbook.create_sheet("Extra")
        workbook.save(self.xlsx)
        workbook.close()


class Float32Tests(unittest.TestCase):
    def test_known_vectors_both_word_orders(self):
        vectors = [
            (0.0, [0, 0]), (-0.0, [0x8000, 0]),
            (1.0, [0x3F80, 0]), (-2.5, [0xC020, 0]),
            (123.25, [0x42F6, 0x8000]),
            (2 ** -149, [0, 1]),
            (struct.unpack(">f", bytes.fromhex("7f7fffff"))[0], [0x7F7F, 0xFFFF]),
        ]
        for value, words in vectors:
            for order, expected in (("big", words), ("little", words[::-1])):
                with self.subTest(value=value, order=order):
                    self.assertEqual(formats.encode_float32(value, order), expected)
                    decoded = formats.decode_float32(expected, order)
                    self.assertEqual(decoded, value)
                    self.assertEqual(math.copysign(1, decoded), math.copysign(1, value))

    def test_fraction_rounds_to_float32(self):
        words = formats.encode_float32(0.1)
        self.assertEqual(words, [0x3DCC, 0xCCCD])
        self.assertEqual(formats.decode_float32(words), struct.unpack(">f", struct.pack(">f", 0.1))[0])

    def test_decode_retains_nonfinite_ieee_values(self):
        self.assertTrue(math.isnan(formats.decode_float32([0x7FC0, 0])))
        self.assertEqual(formats.decode_float32([0x7F80, 0]), math.inf)
        self.assertEqual(formats.decode_float32([0xFF80, 0]), -math.inf)

    def test_encode_rejects_nonfinite_overflow_and_wrong_types(self):
        for value in (math.inf, -math.inf, math.nan, 3.5e38, -3.5e38,
                      10 ** 1000, True, False, "1", None, [], complex(1, 0)):
            with self.subTest(value=repr(value)):
                with self.assertRaises(ValueError):
                    formats.encode_float32(value)

    def test_decode_requires_exactly_two_uint16_words(self):
        for words in ([], [1], [1, 2, 3], (1, 2), "12", None,
                      [True, 0], [0, False], [0.0, 0], [-1, 0],
                      [0, 65536], [0, "1"], [0, math.inf]):
            with self.subTest(words=words):
                with self.assertRaises(ValueError):
                    formats.decode_float32(words)

    def test_order_is_strict_on_both_apis(self):
        for order in ("BIG", "Little", "ABCD", "", None, 1, [], True):
            with self.subTest(order=order):
                with self.assertRaises(ValueError):
                    formats.decode_float32([0, 0], order)
                with self.assertRaises(ValueError):
                    formats.encode_float32(0.0, order)


class XLSXTests(LocalFilesTestCase):
    def test_all_object_types_round_trip_with_exact_schema(self):
        rows = [
            ("Coils", 0, 1), ("Discrete inputs", 65535, 0),
            ("Holding registers", 500, 65535), ("Input registers", 1000, 1234),
        ]
        formats.save_xlsx(self.xlsx, iter(rows))
        self.assertEqual(formats.load_xlsx(self.xlsx), rows)
        workbook = load_workbook(self.xlsx, data_only=False)
        self.assertEqual(tuple(cell.value for cell in workbook.active[1]), formats.XLSX_FIELDS)
        self.assertEqual(workbook.active["B2"].data_type, "n")
        workbook.close()

    def test_empty_export_has_header_and_loads_empty(self):
        formats.save_xlsx(self.xlsx, [])
        self.assertEqual(formats.load_xlsx(self.xlsx), [])

    def test_integer_valued_excel_numbers_are_accepted(self):
        self.workbook([formats.XLSX_FIELDS, ("Coils", 0.0, 1.0)])
        self.assertEqual(formats.load_xlsx(self.xlsx), [("Coils", 0, 1)])

    def test_strict_headers(self):
        for header in ((), ("address", "object_type", "value"),
                       ("object_type", "address"), (*formats.XLSX_FIELDS, "extra"),
                       (" object_type", "address", "value"),
                       ("=1+1", "address", "value")):
            with self.subTest(header=header):
                self.workbook([header])
                with self.assertRaises(ValueError):
                    formats.load_xlsx(self.xlsx)

    def test_malformed_untrusted_rows_are_rejected(self):
        rows = [
            ("Coils", 0), ("Coils", 0, 1, "extra"),
            ("Coils", 0, 1, None), ("Unknown", 0, 1),
            (" Coils", 0, 1), ("Coils", -1, 1), ("Coils", 65536, 1),
            ("Coils", 0, 2), ("Discrete inputs", 0, -1),
            ("Holding registers", 0, 65536), ("Input registers", 0, -1),
            ("Input registers", 1.5, 1), ("Input registers", 0, 1.5),
            ("Coils", "0", 1), ("Coils", 0, "1"),
            ("Holding registers", "0x100", 1),
            ("Coils", True, 1), ("Coils", 0, False),
            ("Coils", None, 1), ("Coils", 0, None),
            ("Coils", "=0", 1), ("Coils", 0, "=1"),
            ("Coils", 0, "#VALUE!"), (None, None, None),
            ("Coils", 0, math.inf), ("Coils", math.nan, 1),
        ]
        for row in rows:
            with self.subTest(row=row):
                self.workbook([formats.XLSX_FIELDS, ("Coils", 0, 1), row])
                with self.assertRaises(ValueError):
                    formats.load_xlsx(self.xlsx)

    def test_extra_sheets_and_merged_missing_cells_rejected(self):
        self.workbook([formats.XLSX_FIELDS], extra_sheet=True)
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)
        workbook = Workbook()
        workbook.active.append(formats.XLSX_FIELDS)
        workbook.active.append(("Coils", 0, 1))
        workbook.active.merge_cells("B2:C2")
        workbook.save(self.xlsx)
        workbook.close()
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)

    def test_formula_with_cached_value_is_rejected(self):
        self.workbook([formats.XLSX_FIELDS, ("Coils", 0, "=1")])
        with ZipFile(self.xlsx) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        root = ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        root.find(".//s:c[@r='C2']/s:v", ns).text = "1"
        entries["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
        with ZipFile(self.xlsx, "w") as target:
            for name, data in entries.items():
                target.writestr(name, data)
        cached = load_workbook(self.xlsx, data_only=True)
        self.assertEqual(cached.active["C2"].value, 1)
        cached.close()
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)

    def test_misleading_dimensions_do_not_hide_invalid_rows(self):
        self.workbook([formats.XLSX_FIELDS, ("Coils", 0, 2)])
        with ZipFile(self.xlsx) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
            b'ref="A1:C2"', b'ref="A1:C1"',
        )
        with ZipFile(self.xlsx, "w") as target:
            for name, data in entries.items():
                target.writestr(name, data)
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)

    def test_invalid_export_never_overwrites_destination(self):
        self.xlsx.write_bytes(b"original")
        for row in (("Coils", 0, "=1"), ("Coils", 0, True),
                    ("Holding registers", 0, 1.0), {"object_type": "Coils"},
                    ("Coils", 0, 1, "extra")):
            with self.subTest(row=row):
                with self.assertRaises(ValueError):
                    formats.save_xlsx(self.xlsx, [("Coils", 0, 1), row])
                self.assertEqual(self.xlsx.read_bytes(), b"original")

    def test_actual_row_limit(self):
        self.assertEqual(formats.MAX_ROWS, 100_000)
        rows = [("Coils", 0, 1)] * formats.MAX_ROWS
        formats.save_xlsx(self.xlsx, rows)
        self.assertEqual(len(formats.load_xlsx(self.xlsx)), formats.MAX_ROWS)
        before = self.xlsx.stat().st_size
        with self.assertRaisesRegex(ValueError, "100,000"):
            formats.save_xlsx(self.xlsx, rows + [rows[0]])
        self.assertEqual(self.xlsx.stat().st_size, before)
        with patch.object(formats, "MAX_ROWS", 99_999):
            with self.assertRaisesRegex(ValueError, "100,000"):
                formats.load_xlsx(self.xlsx)

    def test_invalid_zip_and_missing_file(self):
        self.xlsx.write_bytes(b"not an XLSX file")
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)
        self.xlsx.unlink()
        with self.assertRaises(FileNotFoundError):
            formats.load_xlsx(self.xlsx)

    def test_noniterable_export_and_invalid_xml_are_rejected(self):
        for rows in (None, 42, True):
            with self.assertRaises(ValueError):
                formats.save_xlsx(self.xlsx, rows)
        self.workbook([formats.XLSX_FIELDS])
        with ZipFile(self.xlsx) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        entries["xl/worksheets/sheet1.xml"] = b"<worksheet><unclosed>"
        with ZipFile(self.xlsx, "w") as target:
            for name, data in entries.items():
                target.writestr(name, data)
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)

    def test_falsely_narrow_dimensions_do_not_hide_extra_columns(self):
        self.workbook([formats.XLSX_FIELDS, ("Coils", 0, 1, "hidden extra")])
        with ZipFile(self.xlsx) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
            b'ref="A1:D2"', b'ref="A1:C2"',
        )
        with ZipFile(self.xlsx, "w") as target:
            for name, data in entries.items():
                target.writestr(name, data)
        with self.assertRaises(ValueError):
            formats.load_xlsx(self.xlsx)


class ProfileTests(LocalFilesTestCase):
    def write_document(self, document):
        self.profile.write_text(json.dumps(document), encoding="utf-8")

    def document(self, prefs=None):
        return {"format": "m31tool-profile", "version": 1,
                "settings": settings() if prefs is None else prefs}

    def test_tcp_and_rtu_round_trip_canonical_and_no_mutation(self):
        for prefs in (settings(host="localhost"),
                      settings(transport="RTU", serial_port="/dev/ttyUSB2",
                               parity="E", stopbits=2, unit_id=247)):
            with self.subTest(transport=prefs["transport"]):
                before = prefs.copy()
                self.assertIsNone(formats.save_profile(self.profile, prefs))
                self.assertEqual(prefs, before)
                self.assertEqual(formats.load_profile(self.profile), prefs)
                self.assertEqual(json.loads(self.profile.read_text()), self.document(prefs))

    def test_blank_endpoints_are_preserved_without_hardware_defaults(self):
        formats.save_profile(self.profile, settings())
        prefs = formats.load_profile(self.profile)
        self.assertEqual(prefs["host"], "")
        self.assertEqual(prefs["serial_port"], "")
        self.assertEqual(set(prefs), set(formats.PROFILE_FIELDS))

    def test_missing_and_unknown_fields_on_save_and_load(self):
        variants = [dict(settings(), password="secret"), dict(settings(), write_value=1),
                    dict(settings(), device_configuration={}), [], None]
        for field in formats.PROFILE_FIELDS:
            prefs = settings()
            del prefs[field]
            variants.append(prefs)
        for prefs in variants:
            with self.subTest(prefs=prefs):
                with self.assertRaises(ValueError):
                    formats.save_profile(self.profile, prefs)
                self.write_document({"format": "m31tool-profile", "version": 1, "settings": prefs})
                with self.assertRaises(ValueError):
                    formats.load_profile(self.profile)

    def test_all_invalid_settings_on_save_and_load(self):
        invalid = {
            "transport": ["tcp", "", None, 1],
            "host": [1, None, " host ", "a\nb", "a\x00b", "a\x7fb", "h" * 1025],
            "serial_port": [1, "/dev/x\n", " /dev/x"],
            "tcp_port": [0, 65536, True, 502.0, "502"],
            "baudrate": [0, 4_000_001, False, 9600.0, "9600"],
            "parity": ["n", "X", 1],
            "stopbits": [0, 3, 1.5, True],
            "unit_id": [0, 248, -1, True, "1"],
            "timeout": [0, -1, 60.1, math.inf, math.nan, True, "2", 10 ** 1000],
            "object_type": ["Unknown", " Holding registers", None, []],
            "address": [-1, 65536, True, 0.0, "0"],
            "count": [0, 126, 2001, True, 1.0, "1"],
        }
        for field, values in invalid.items():
            for value in values:
                with self.subTest(field=field, value=repr(value)):
                    prefs = settings(**{field: value})
                    with self.assertRaises(ValueError):
                        formats.save_profile(self.profile, prefs)
                    self.write_document(self.document(prefs))
                    with self.assertRaises(ValueError):
                        formats.load_profile(self.profile)

    def test_count_limits_and_address_end_inclusive(self):
        for kind, maximum in zip(formats.OBJECT_TYPES, (2000, 2000, 125, 125)):
            prefs = settings(object_type=kind, count=maximum, address=65536 - maximum,
                             timeout=60, tcp_port=65535, baudrate=4_000_000)
            formats.save_profile(self.profile, prefs)
            self.assertEqual(formats.load_profile(self.profile), prefs)
            for invalid in (dict(prefs, count=maximum + 1), dict(prefs, address=prefs["address"] + 1)):
                with self.assertRaises(ValueError):
                    formats.save_profile(self.profile, invalid)
        formats.save_profile(self.profile, settings(address=65535, count=1))

    def test_top_level_and_version_are_strict(self):
        documents = [None, [], {}, {"settings": settings()},
                     dict(self.document(), extra="unknown")]
        for value in (0, 2, "1", 1.0, True, None):
            documents.append(dict(self.document(), version=value))
        for value in ("other", None, 1, []):
            documents.append(dict(self.document(), format=value))
        for document in documents:
            with self.subTest(document=document):
                self.write_document(document)
                with self.assertRaises(ValueError):
                    formats.load_profile(self.profile)

    def test_duplicate_fields_invalid_json_and_nonfinite_numbers(self):
        valid = json.dumps(self.document())
        documents = [
            "", "{", valid.replace('"version": 1', '"version": 1, "version": 1'),
            valid.replace('"unit_id": 1', '"unit_id": 1, "unit_id": 2'),
            valid.replace('"timeout": 2.0', '"timeout": NaN'),
            valid.replace('"timeout": 2.0', '"timeout": Infinity'),
            valid.replace('"timeout": 2.0', '"timeout": -Infinity'),
            valid.replace('"timeout": 2.0', '"timeout": 1e9999'),
        ]
        for content in documents:
            with self.subTest(content=content):
                self.profile.write_text(content, encoding="utf-8")
                with self.assertRaises(ValueError):
                    formats.load_profile(self.profile)

    def test_size_limit_and_invalid_utf8(self):
        self.profile.write_bytes(b" " * 65_537)
        with self.assertRaisesRegex(ValueError, "65,536"):
            formats.load_profile(self.profile)
        self.profile.write_bytes(b"\xff")
        with self.assertRaises(ValueError):
            formats.load_profile(self.profile)

    def test_excessive_json_nesting_is_a_validation_error(self):
        self.profile.write_text("[" * 2000 + "]" * 2000, encoding="utf-8")
        with self.assertRaises(ValueError):
            formats.load_profile(self.profile)

    def test_invalid_save_does_not_overwrite_existing_file(self):
        self.profile.write_text("original", encoding="utf-8")
        with self.assertRaises(ValueError):
            formats.save_profile(self.profile, settings(unit_id=0))
        self.assertEqual(self.profile.read_text(), "original")

    def test_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            formats.load_profile(self.profile)

    def test_module_and_every_api_are_local_only_in_fresh_interpreter(self):
        script = """
import importlib.abc
import json
from pathlib import Path
import socket
import sys
from unittest.mock import patch
class Guard(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname in ('m31tool.app', 'm31tool.core', 'pymodbus', 'serial', 'tkinter'):
            raise AssertionError('forbidden import: ' + fullname)
sys.meta_path.insert(0, Guard())
with patch.object(socket, 'socket') as sock, patch.object(socket, 'create_connection') as conn:
    from m31tool.formats import decode_float32, encode_float32, save_xlsx, load_xlsx, save_profile, load_profile
    directory = Path(sys.argv[1])
    assert decode_float32(encode_float32(1.0)) == 1.0
    save_xlsx(directory / 'local.xlsx', [('Coils', 0, 1)])
    assert load_xlsx(directory / 'local.xlsx') == [('Coils', 0, 1)]
    prefs = json.loads(sys.argv[2])
    save_profile(directory / 'local.json', prefs)
    assert load_profile(directory / 'local.json') == prefs
    sock.assert_not_called()
    conn.assert_not_called()
"""
        result = subprocess.run(
            [sys.executable, "-c", script, str(self.directory), json.dumps(settings())],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
