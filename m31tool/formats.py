"""Validated local display formats and client preferences; never device I/O.

XLSX files contain one worksheet with exactly ``object_type,address,value``
columns and numeric address/value cells. Profiles require every field in
PROFILE_FIELDS: no endpoint or hardware defaults are inferred. Blank endpoints
are allowed for disconnected client preferences. Profiles contain no write
values or device configuration and loading never applies settings.
"""

from __future__ import annotations

import json
import math
import re
import struct
from contextlib import ExitStack
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException


OBJECT_TYPES = (
    "Coils", "Discrete inputs", "Holding registers", "Input registers",
)
XLSX_FIELDS = ("object_type", "address", "value")
MAX_ROWS = 100_000
PROFILE_FIELDS = (
    "transport", "host", "tcp_port", "serial_port", "baudrate", "parity",
    "stopbits", "unit_id", "timeout", "object_type", "address", "count",
)
_MAX_PROFILE_BYTES = 65_536
_MAX_FLOAT32_WRITE_VALUES = 61
_FLOAT32_LITERAL = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?")


def _integer(value, name: str, minimum: int, maximum: int) -> int:
    if type(value) is not int or not minimum <= value <= maximum:
        raise ValueError(f"{name} must be an integer between {minimum} and {maximum}.")
    return value


def _choice(value, name: str, choices) -> str:
    if type(value) is not str or value not in choices:
        raise ValueError(f"{name} must be one of {', '.join(choices)}.")
    return value


def _word_order(word_order: str) -> str:
    return _choice(word_order, "word_order", ("big", "little"))


def decode_float32(words: list[int], word_order: str = "big") -> float:
    """Decode two uint16 words; big = high word first, bytes always big-endian.

    IEEE nonfinite values are retained on decode so callers can display them.
    """
    order = _word_order(word_order)
    if type(words) is not list or len(words) != 2:
        raise ValueError("float32 requires a list of exactly two words.")
    for word in words:
        _integer(word, "word", 0, 65535)
    high, low = words if order == "big" else words[::-1]
    return struct.unpack(">f", struct.pack(">HH", high, low))[0]


def encode_float32(value: float, word_order: str = "big") -> list[int]:
    """Encode a finite float32, rejecting overflow rather than writing infinity."""
    order = _word_order(word_order)
    if type(value) not in (int, float):
        raise ValueError("float32 value must be a finite number.")
    try:
        if not math.isfinite(value):
            raise ValueError("float32 value must be finite.")
        words = list(struct.unpack(">HH", struct.pack(">f", value)))
    except (OverflowError, struct.error) as exc:
        raise ValueError("value is outside the finite float32 range.") from exc
    return words if order == "big" else words[::-1]


def parse_float32_write(text: str, word_order: str) -> list[int]:
    """Validate 1–61 decimal float literals and return holding-register words.

    Commas separate values, never decimal fractions; surrounding whitespace is
    allowed. Explicit big/little word order is required, with big-endian bytes
    within each word. Nonfinite values and float32 overflow are rejected.
    No I/O occurs and no partial result is returned on invalid input.
    """
    order = _word_order(word_order)
    if type(text) is not str:
        raise ValueError("float32 write text must be a comma-separated string.")
    tokens = text.split(",", _MAX_FLOAT32_WRITE_VALUES)
    if len(tokens) > _MAX_FLOAT32_WRITE_VALUES:
        raise ValueError("float32 writes require at most 61 values (122 registers).")
    words: list[int] = []
    for token in tokens:
        literal = token.strip()
        if _FLOAT32_LITERAL.fullmatch(literal) is None:
            raise ValueError("float32 writes require nonempty decimal float literals.")
        words.extend(encode_float32(float(literal), order))
    return words


def _row(row, line: int, *, excel: bool = False) -> tuple[str, int, int]:
    if not isinstance(row, (tuple, list)) or len(row) != 3:
        raise ValueError(f"Row {line} must have exactly three columns.")
    kind = _choice(row[0], f"Row {line} object_type", OBJECT_TYPES)
    numbers = []
    for name, raw, maximum in (
        ("address", row[1], 65535),
        ("value", row[2], 1 if kind in OBJECT_TYPES[:2] else 65535),
    ):
        # Excel represents numbers as either ints or floats, not typed uint16s.
        if excel and type(raw) is float and math.isfinite(raw) and raw.is_integer():
            raw = int(raw)
        numbers.append(_integer(raw, f"Row {line} {name}", 0, maximum))
    return kind, numbers[0], numbers[1]


def load_xlsx(path) -> list[tuple[str, int, int]]:
    """Load a strictly validated workbook without evaluating formulas or links."""
    workbook = None
    iterator = None
    resources = ExitStack()
    try:
        # Own the file even when openpyxl raises before returning a workbook.
        source = resources.enter_context(open(path, "rb"))
        workbook = load_workbook(
            source, read_only=True, data_only=False, keep_links=False,
        )
        if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1:
            raise ValueError("XLSX must contain exactly one worksheet.")
        sheet = workbook.worksheets[0]
        if sheet.max_column is not None and sheet.max_column > len(XLSX_FIELDS):
            raise ValueError("XLSX has extra columns.")
        # Do not trust dimensions supplied by an external workbook producer.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        if (tuple(cell.value for cell in header) != XLSX_FIELDS
                or any(cell.data_type != "s" for cell in header)):
            raise ValueError("XLSX header must be exactly: object_type,address,value")
        rows = []
        for line, cells in enumerate(iterator, 2):
            if len(rows) >= MAX_ROWS:
                raise ValueError("XLSX exceeds the 100,000-row local import limit.")
            if len(cells) != 3 or any(cell.data_type in ("f", "e") for cell in cells):
                raise ValueError(f"XLSX row {line} has formulas, errors or extra/missing columns.")
            rows.append(_row(tuple(cell.value for cell in cells), line, excel=True))
        return rows
    except (InvalidFileException, BadZipFile, SyntaxError, KeyError) as exc:
        raise ValueError("Invalid XLSX workbook.") from exc
    finally:
        try:
            if iterator is not None:
                iterator.close()
        finally:
            try:
                if workbook is not None:
                    workbook.close()
            finally:
                resources.close()


def save_xlsx(path, rows) -> None:
    """Validate all typed rows before opening the local destination for export."""
    try:
        iterator = iter(rows)
    except TypeError as exc:
        raise ValueError("XLSX rows must be an iterable of three-column rows.") from exc
    validated = []
    for line, row in enumerate(iterator, 2):
        if len(validated) >= MAX_ROWS:
            raise ValueError("XLSX exceeds the 100,000-row local export limit.")
        validated.append(_row(row, line))
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Values"
        sheet.append(XLSX_FIELDS)
        for row in validated:
            sheet.append(row)
        workbook.save(path)
    finally:
        workbook.close()


def _settings(settings: dict) -> dict:
    if type(settings) is not dict or set(settings) != set(PROFILE_FIELDS):
        raise ValueError("Profile settings must contain exactly: " + ", ".join(PROFILE_FIELDS))
    result = {field: settings[field] for field in PROFILE_FIELDS}
    _choice(result["transport"], "transport", ("TCP", "RTU"))
    for field in ("host", "serial_port"):
        value = result[field]
        if (type(value) is not str or len(value) > 1024
                or value != value.strip()
                or any(ord(char) < 32 or ord(char) == 127 for char in value)):
            raise ValueError(f"{field} must be a trimmed string without control characters (max 1024).")
    for field, minimum, maximum in (
        ("tcp_port", 1, 65535), ("baudrate", 1, 4_000_000),
        ("stopbits", 1, 2), ("unit_id", 1, 247), ("address", 0, 65535),
    ):
        _integer(result[field], field, minimum, maximum)
    _choice(result["parity"], "parity", ("N", "E", "O"))
    timeout = result["timeout"]
    if type(timeout) not in (int, float) or not 0 < timeout <= 60:
        raise ValueError("timeout must be a finite number greater than zero and at most 60.")
    result["timeout"] = float(timeout)
    kind = _choice(result["object_type"], "object_type", OBJECT_TYPES)
    _integer(result["count"], "count", 1, 2000 if kind in OBJECT_TYPES[:2] else 125)
    if result["address"] + result["count"] > 65536:
        raise ValueError("address + count exceeds the Modbus address space.")
    return result


def save_profile(path, settings: dict) -> None:
    """Save client/read preferences only, with no defaults or device operations."""
    document = {"format": "m31tool-profile", "version": 1, "settings": _settings(settings)}
    content = json.dumps(document, indent=2, ensure_ascii=True, allow_nan=False) + "\n"
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(content)


def _unique_object(pairs) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON field: {key}")
        result[key] = value
    return result


def _reject_constant(value):
    raise ValueError(f"Nonfinite JSON number: {value}")


def load_profile(path) -> dict:
    """Return validated preferences; never connect, write or apply configuration."""
    with open(path, "rb") as handle:
        content = handle.read(_MAX_PROFILE_BYTES + 1)
    if len(content) > _MAX_PROFILE_BYTES:
        raise ValueError("Profile exceeds the 65,536-byte local import limit.")
    try:
        document = json.loads(
            content.decode("utf-8"), object_pairs_hook=_unique_object,
            parse_constant=_reject_constant,
        )
    except RecursionError as exc:
        raise ValueError("Profile JSON is nested too deeply.") from exc
    if type(document) is not dict or set(document) != {"format", "version", "settings"}:
        raise ValueError("Profile must contain exactly format, version and settings.")
    if document["format"] != "m31tool-profile":
        raise ValueError("Unsupported profile format.")
    if type(document["version"]) is not int or document["version"] != 1:
        raise ValueError("Unsupported profile version.")
    return _settings(document["settings"])
